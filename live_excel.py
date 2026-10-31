import os
import openpyxl
import xlwings as xw

WORKBOOK_NAME = "Live_Nifty_Tracker.xlsx"
WORKBOOK_PATH = os.path.abspath(WORKBOOK_NAME)

_app_instance = None
_wb_instance = None

def _ensure_workbook_file():
    """Ensure the physical workbook file exists with initial sheets."""
    if not os.path.exists(WORKBOOK_PATH):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Option Chain"
        wb.create_sheet(title="Nifty 50 Stocks")
        wb.save(WORKBOOK_PATH)

def get_or_create_sheet(sheet_name="Option Chain"):
    """
    Connect to Live_Nifty_Tracker.xlsx or launch Excel and open it.
    Reuses persistent app and workbook handles across Streamlit reruns.
    """
    global _app_instance, _wb_instance
    _ensure_workbook_file()

    need_connect = True
    if _wb_instance is not None and _app_instance is not None:
        try:
            _ = _wb_instance.name
            need_connect = False
        except Exception:
            _app_instance = None
            _wb_instance = None
            need_connect = True

    if need_connect:
        for app in xw.apps:
            for b in app.books:
                if b.name.lower() == WORKBOOK_NAME.lower():
                    _app_instance = app
                    _wb_instance = b
                    need_connect = False
                    break
            if not need_connect:
                break

    if need_connect:
        _app_instance = xw.App(visible=True, add_book=False)
        _wb_instance = _app_instance.books.open(WORKBOOK_PATH)

    try:
        sheet = _wb_instance.sheets[sheet_name]
    except Exception:
        sheet = _wb_instance.sheets.add(sheet_name)

    return sheet

def stream_to_excel(df, metrics, title="NIFTY Live Tracker", sheet_name="Option Chain"):
    """
    Stream a pandas DataFrame and headline metrics into Live_Nifty_Tracker.xlsx in real-time.
    Also calls book.save() so both the live open Excel screen and the file stay synchronized.
    """
    sheet = get_or_create_sheet(sheet_name)

    # 1. Update Header & Summary Metrics at the top
    sheet.range("A1").value = title
    sheet.range("A1").font.bold = True
    sheet.range("A1").font.size = 14

    underlying = metrics.get("underlying", "")
    atm = metrics.get("atm", "")
    pcr = round(metrics.get("pcr_overall", 0), 2)

    sheet.range("A2").value = f"Underlying LTP: {underlying}"
    if atm:
        sheet.range("A3").value = f"ATM: {atm}   PCR: {pcr}"
    else:
        sheet.range("A3").value = ""

    # 2. Filter out internal columns
    cols = [c for c in df.columns if not c.startswith("_")]
    export_df = df[cols].copy()

    # 3. Write Headers and Table values in one high-performance block
    start_row = 5
    headers_range = sheet.range((start_row, 1), (start_row, len(cols)))
    headers_range.value = cols
    headers_range.font.bold = True
    headers_range.color = (220, 220, 220)

    data_range = sheet.range((start_row + 1, 1), (start_row + len(export_df), len(cols)))
    data_range.value = export_df.values

    # 4. Apply Percentage format for Chg% columns
    for j, c in enumerate(cols, start=1):
        if "Chg%" in c:
            pct_range = sheet.range((start_row + 1, j), (start_row + len(export_df), j))
            pct_range.number_format = "0.00%"

    # 5. Save workbook so disk file and live view are in sync
    try:
        sheet.book.save()
    except Exception:
        pass

    return True
